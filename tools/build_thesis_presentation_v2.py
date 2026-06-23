"""Generate thesis presentation PPTX (V2) from structured content.

V2 vs V1:
- Domain-general introduction & problem statement (not bound to the two base papers)
- Literature review expanded to 6+ recent works (split across 3 slides)
- Base-paper-style diagram slides described in text (export Mermaid from
  doc/THESIS_PRESENTATION_V2.md for the actual figures)

Outputs doc/THESIS_PRESENTATION_V2.pptx. Does NOT touch the V1 files.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "THESIS_PRESENTATION_V2.pptx"


def _set_title(slide, text: str, subtitle: str | None = None):
    slide.shapes.title.text = text
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def _add_bullets(slide, lines, left=0.5, top=1.4, width=9.0, height=5.5, font_size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        line = line.strip().replace("**", "")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if line.startswith("- "):
            line = line[2:]
        p.text = line
        p.level = 0
        p.font.size = Pt(font_size)


def _add_table(slide, headers, rows, top=1.5, col_widths=None):
    nrows = len(rows) + 1
    ncols = len(headers)
    height = Inches(0.35 * nrows)
    table = slide.shapes.add_table(nrows, ncols, Inches(0.4), Inches(top), Inches(9.2), height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(11)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val).replace("**", "")
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)


def _add_flow_text(slide, title: str, lines, top=1.3):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(0.4))
    box.text_frame.paragraphs[0].text = title
    box.text_frame.paragraphs[0].font.bold = True
    box.text_frame.paragraphs[0].font.size = Pt(14)
    _add_bullets(slide, lines, top=top + 0.35, font_size=14)


def _note(slide, text: str):
    """Add a small note pointing to the Mermaid figure to export."""
    box = slide.shapes.add_textbox(Inches(0.5), Inches(6.7), Inches(9), Inches(0.5))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.italic = True


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    # 1 — Title
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = "Autoregressive Chunk-Wise Editing for\nLong-Form Knowledge in Vision-Language Models"
    s.placeholders[1].text = (
        "A study on the VLKEB-long benchmark (BLIP2-OPT-2.7b)\n"
        "MSc Thesis Presentation\n"
        "[Your Name] - [Supervisor] - [Institution] - June 2026"
    )

    # 2 — Introduction (domain-general)
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Introduction")
    _add_bullets(s, [
        "Large models memorize facts in pre-training, but facts become stale, wrong, or harmful over time.",
        "Retraining to fix them is expensive and risky (cost, carbon, regression).",
        "Knowledge editing updates a specific fact while leaving everything else intact.",
        "Good edits satisfy: Reliability (fact updated), Generality (paraphrases), Locality (others preserved).",
        "Two open frontiers: (1) multimodal (vision-language) editing, (2) long-form (sentence/paragraph) edits.",
    ], top=1.25, font_size=17)

    # 3 — Problem statement (domain-general)
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Problem Statement")
    _add_bullets(s, [
        "Editing methods & benchmarks are dominated by short, triplet-style targets.",
        "1) Efficacy degrades with length - few hidden-state edits lose grip on distant tokens.",
        "2) Evaluation breaks - exact match is near-zero; need token/chunk/semantic metrics.",
        "3) Lifelong interference - many sequential edits erode locality and earlier edits.",
        "RQ1: Can autoregressive chunk-wise editing improve long-form VLM edits, keeping generality & locality?",
        "RQ2: How does editing granularity (chunk size) affect long-form performance?",
    ], top=1.2, font_size=16)

    # 4 — Aim, objectives, contributions
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Aim, Objectives & Contributions")
    _add_table(s,
        ["#", "Objective", "Status"],
        [
            ["1", "Construct long-form multimodal benchmark (VLKEB-long)", "Done"],
            ["2", "Implement chunk-wise autoregressive editing", "Done"],
            ["3", "Train single-pass (baseline) vs AR variants", "Done"],
            ["4", "Full-scale eval (~3150 sequential edits)", "Done"],
            ["5", "Granularity ablation (chunk size 8/16/32/64)", "Done"],
        ],
        top=1.35, col_widths=[0.5, 7.0, 1.4])
    _add_bullets(s, [
        "Contributions: long-form multimodal benchmark + protocol; single-pass vs chunk-wise study; granularity analysis.",
    ], top=4.3, font_size=14)

    # 5 — Background concept (diagram)
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Background - What is a Knowledge Edit?")
    _add_flow_text(s, "An edit changes the target response; locality probes stay unchanged:", [
        "Pre-edit:  Q 'Who painted this?' -> wrong artist",
        "Edit:      (image, prompt) -> new target",
        "Post-edit: edited query -> correct;  unrelated query -> unchanged",
        "Three properties scored: Reliability / Generality / Locality.",
    ])
    _note(s, "Figure: see Slide 5 Mermaid in THESIS_PRESENTATION_V2.md")

    # 6 — Lit review map
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review - Landscape")
    _add_bullets(s, [
        "Text LLM editing: locate-then-edit (ROME, MEMIT), hypernetworks (MEND).",
        "Lifelong / sequential editing: WikiBigEdit (500K+), ENCORE (10K edits).",
        "Long-form / unstructured text: UnKE, AnyEdit (chunk-wise autoregressive).",
        "Multimodal editing: MMEdit -> VLKEB -> MMKE-Bench / MC-MKE -> LiveEdit.",
        "This thesis sits at long-form + multimodal + lifelong (under-explored).",
    ], top=1.3, font_size=16)
    _note(s, "Figure: see Slide 6 mindmap in THESIS_PRESENTATION_V2.md")

    # 7 — Lit review 1: foundations
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review (1/3) - Foundations & Surveys")
    _add_table(s,
        ["Work", "Year", "Contribution"],
        [
            ["KE Survey (Wang et al.)", "2023-24", "Editing as constrained opt: accuracy/locality/generality"],
            ["Pitfalls of KE (survey)", "2024", "Editing distorts knowledge & degrades general ability"],
            ["Dual-Axis Taxonomy", "2025", "Function-based view; effect depends on knowledge type"],
        ],
        top=1.4, col_widths=[2.6, 1.1, 5.5])
    _add_bullets(s, [
        "Takeaway: editing is well-formalized for short factual knowledge; long-form/multimodal are open.",
    ], top=4.2, font_size=14)

    # 8 — Lit review 2: lifelong & long-form
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review (2/3) - Lifelong & Long-Form")
    _add_table(s,
        ["Work", "Year", "Contribution"],
        [
            ["WikiBigEdit", "2025", "500K+ real Wikidata edits; lifelong at scale"],
            ["ENCORE", "2025", "Norm-constrained; 10,000 edits without degradation"],
            ["UnKE", "2025", "Edits unstructured / long knowledge in text"],
            ["AnyEdit (base)", "2025", "Autoregressive chunk-wise editing; ROUGE-L/BERTScore"],
        ],
        top=1.4, col_widths=[2.4, 1.1, 5.7])
    _add_bullets(s, [
        "Takeaway: chunk-wise / AR editing solves length in TEXT; not yet studied for vision-language models.",
    ], top=4.4, font_size=14)

    # 9 — Lit review 3: multimodal
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review (3/3) - Multimodal Editing")
    _add_table(s,
        ["Work", "Year", "Contribution"],
        [
            ["MMEdit", "2023", "First multimodal editing benchmark"],
            ["VLKEB", "2024", "Real images via KG; adds Portability; harder locality"],
            ["MMKE-Bench", "2025", "Free-form visual knowledge; entity/semantic/user edits"],
            ["MC-MKE", "2025", "Fine-grained edits; modality consistency"],
            ["LiveEdit (base)", "2025", "Lifelong VLM editing via low-rank mixture-of-experts"],
        ],
        top=1.4, col_widths=[2.2, 1.1, 5.9])
    _add_bullets(s, [
        "Gap: no prior work studies long-form targets in a lifelong VLM editing setting.",
    ], top=4.5, font_size=14)

    # 10 — Research gap quadrant
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Research Gap & Positioning")
    _add_table(s,
        ["Work", "Modality", "Lifelong", "Long targets"],
        [
            ["ROME / MEMIT / MEND", "LLM", "Limited", "Poor"],
            ["UnKE / AnyEdit", "LLM", "Batch", "Strong"],
            ["MMEdit / VLKEB / LiveEdit", "VLM", "Yes", "Short"],
            ["This thesis", "VLM", "Yes (seq=50)", "Long (VLKEB-long)"],
        ],
        top=1.4, col_widths=[3.0, 1.4, 1.6, 2.2])
    _add_bullets(s, [
        "Target quadrant: long-form + multimodal + lifelong.",
    ], top=4.4, font_size=14)
    _note(s, "Figure: see Slide 10 quadrantChart in THESIS_PRESENTATION_V2.md")

    # 11 — LiveEdit Fig.1-style concept
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Concept - Lifelong VLM Editing (base-paper style)")
    _add_flow_text(s, "Stream of edits over time; model must stay correct AND stable:", [
        "Edit stream: edit_1 -> edit_2 -> ... -> edit_t",
        "Reliability: edited query -> new answer",
        "Generality: rephrased text / image -> new answer",
        "Locality: unrelated text / image -> unchanged",
    ])
    _note(s, "Recreates LiveEdit Fig.1 - export Slide 11 Mermaid or screenshot the PDF figure")

    # 12 — LiveEdit Fig.2-style architecture
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Architecture - External MoE Editor (base-paper style)")
    _add_flow_text(s, "Per-edit experts + routing over a frozen VLM:", [
        "Edit sample (v_e, p_e, o_e) -> expert generator -> low-rank expert (U,V)",
        "Routing features (visual + text) stored with expert in repository",
        "Inference: hard routing (visual filter) -> soft routing (text fusion)",
        "Residual update: h + sum(weighted experts) -> adapted prediction",
        "Editor is external; experts accumulate per edit (lifelong).",
    ])
    _note(s, "Recreates LiveEdit Fig.2 - export Slide 12 Mermaid or screenshot the PDF figure")

    # 13 — AnyEdit Fig.1-style mechanism
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Mechanism - Single-Pass vs Autoregressive (base-paper style)")
    _add_flow_text(s, "Why chunking helps long targets:", [
        "Single-pass: one edit on full ~89-token target -> weak control over distant tokens",
        "Autoregressive: chunk_1 -> chunk_2 -> ... -> chunk_k, each edit conditions on prior chunks",
        "Efficacy vs length: single-pass drops with length; AR stays high (concept).",
    ])
    _note(s, "Recreates AnyEdit Fig.1 - export Slide 13 Mermaid (flow + xychart)")

    # 14 — Methodology pipeline
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Methodology - Overall Pipeline")
    _add_flow_text(s, "End-to-end flow:", [
        "Data: VLKEB (unchanged) -> long-form generator (Ollama) -> VLKEB-long JSON",
        "Train: BLIP2-OPT-2.7b (frozen) + editor; single-pass vs AR (chunk=16)",
        "Eval: sequential edits n=50, seed=42, ~3150 edits",
        "Metrics: Reliability / Generality / Locality / Chunk EM-F1",
    ])
    _note(s, "Figure: see Slide 14 Mermaid in THESIS_PRESENTATION_V2.md")

    # 15 — VLKEB-long construction
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Methodology - VLKEB-long Construction")
    _add_bullets(s, [
        "Official data/images unchanged; only the target is elongated.",
        "alt_short = original short target; long target ~65-150 tokens, anchored to alt_short.",
        "Validation: long target must contain alt_short AND fall in token band.",
        "Scale: 3174 eval rows; 3150 edits; mean length ~89 tokens.",
        "Honest limit: lexical anchoring, not full visual grounding of every clause.",
    ], top=1.3, font_size=16)
    _note(s, "Figure: see Slide 15 Mermaid in THESIS_PRESENTATION_V2.md")

    # 16 — Results main
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Results - Single-Pass vs Autoregressive (cs=16)")
    _add_table(s,
        ["Metric", "Baseline", "AR", "Delta"],
        [
            ["Reliability acc", "75.40%", "75.94%", "+0.54 pp"],
            ["Chunk EM", "13.50%", "14.82%", "+1.32 pp"],
            ["Chunk F1", "76.28%", "76.90%", "+0.62 pp"],
            ["Text generality", "75.21%", "75.63%", "+0.42 pp"],
            ["Image generality", "71.49%", "72.51%", "+1.02 pp"],
            ["Locality (T/I)", "100%/100%", "100%/100%", "0"],
            ["Exact match", "0%", "0%", "-"],
        ],
        top=1.3, col_widths=[2.4, 1.6, 1.6, 1.4])
    _add_bullets(s, [
        "AR improves long-form editing with no locality drop; n=3150, seed 42.",
    ], top=4.9, font_size=14)

    # 17 — Ablation
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Results - Granularity (Chunk Size) Ablation")
    _add_table(s,
        ["chunk_size", "Rel acc", "Chunk EM", "Chunk F1", "Locality"],
        [
            ["8", "75.94%", "26.26%", "76.48%", "100%"],
            ["16 (default)", "75.94%", "14.82%", "76.90%", "100%"],
            ["32", "75.94%", "8.11%", "77.30%", "100%"],
            ["64", "75.94%", "6.14%", "77.88%", "100%"],
        ],
        top=1.4, col_widths=[1.6, 1.3, 1.3, 1.3, 1.2])
    _add_bullets(s, [
        "Token accuracy identical across chunk sizes -> inference granularity doesn't change reliability.",
        "cs=16 matches training; recommended default. Chunk EM not comparable across sizes.",
    ], top=4.5, font_size=14)

    # 18 — Discussion
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Discussion")
    _add_bullets(s, [
        "AR yields consistent, modest gains (rel +0.54 pp, chunk EM +1.32 pp) with no locality cost.",
        "Generality improves on both text and image rephrases.",
        "EM ~ 0% is expected for ~89-token targets -> report token/chunk metrics.",
        "Inference chunk size is secondary when training uses cs=16.",
        "Limits: single backbone (BLIP2), lexical anchoring, no ROUGE-L/BERTScore yet.",
    ], top=1.3, font_size=16)

    # 19 — Limitations & future
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Limitations & Future Work")
    _add_table(s,
        ["Limitation", "Future direction"],
        [
            ["Modest effect size", "Multiple seeds, significance tests"],
            ["Lexical anchoring", "Human audit + grounding-verified tiers"],
            ["Metric set", "Add ROUGE-L, BERTScore (long-form)"],
            ["Mechanism novelty", "Vision-gated chunking; edit-budget; forgetting"],
            ["One backbone", "Add LLaVA / MiniGPT-4"],
            ["Lifelong depth", "Edit-count sweep 1 -> 1000"],
        ],
        top=1.4, col_widths=[3.3, 5.7])

    # 20 — Conclusion
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Conclusion")
    _add_bullets(s, [
        "Identified an open gap: long-form editing in multimodal, lifelong settings.",
        "Built VLKEB-long and a reproducible long-form construction protocol.",
        "Showed autoregressive chunk-wise editing improves long-form reliability & chunk match, no locality loss.",
        "Granularity analysis: cs=16 is a sound default.",
        "Take-home: chunk-wise AR editing is a promising direction for long-form VLM knowledge editing.",
    ], top=1.3, font_size=16)

    # 21 — Q&A
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Thank You - Questions?")
    _add_bullets(s, [
        "Result artifacts: eval_results/comparison/*.json",
        "Base figures for credit: LiveEdit (CVPR 2025) Fig.1-2; AnyEdit (ICML 2025) Fig.1",
        "Anticipated Q: small gains? long outputs + single seed; report token/chunk metrics.",
        "Anticipated Q: benchmark novelty? extension + protocol; anchoring is explicit.",
    ], top=1.4, font_size=16)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
