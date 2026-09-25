"""Generate supervisor progress report DOCX from briefing content + sample images."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
DOC = ROOT / "doc"
SAMPLES = DOC / "figures" / "supervisor_samples"
FIG = DOC / "results" / "meeting_validation" / "step6" / "figures"
OUT = DOC / "THESIS_PROGRESS_REPORT.docx"


def set_run_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True)
    return h


def add_para(doc, text, *, bold=False, italic=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=11)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            run = cells[c_i].paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()


def add_image(doc, path: Path, width=4.2, caption=None):
    if not path.is_file():
        add_para(doc, f"[Missing image: {path.name}]", italic=True, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        r = cap.add_run(caption)
        set_run_font(r, size=9, italic=True, color=RGBColor(0x55, 0x55, 0x55))


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MSc Thesis Progress Report")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Adversarial Visual Binding & Dual-Expert Keys\n"
        "for Lifelong Vision–Language Model Editing"
    )
    set_run_font(r, size=13, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Date: 13 July 2026  ·  Backbone: BLIP2-OPT-2.7B  ·  Base editor: LiveEdit\n"
        "Method: AVBT + Dual-Expert Keys (DEK)"
    )
    set_run_font(r, size=10, italic=True, color=RGBColor(0x44, 0x44, 0x44))

    add_para(
        doc,
        "This report summarises current progress for supervisor discussion: "
        "problem diagnosis, methodology (how and why, with a concrete data sample), "
        "and result analysis. The tone is deliberately honest about what improved "
        "strongly, what improved modestly, and what remains a limitation.",
        size=11,
    )

    # 1 Executive summary
    add_heading(doc, "1. Executive summary", 1)
    add_para(
        doc,
        "Lifelong VLM editors such as LiveEdit can look reliable on the correct image, "
        "but still leak the edited answer when the image is swapped or confusable. "
        "We traced that leakage mainly to false-cross: the system retrieves the wrong "
        "past edit from its lifelong expert pool.",
    )
    add_para(doc, "We propose AVBT + Dual-Expert Keys (DEK):", bold=True)
    add_bullets(
        doc,
        [
            "AVBT — train the editor to bind the right image and suppress adversarial images;",
            "Dual-Expert Keys — at retrieval, fire an edit only if vision and text both agree (vision ∧ text).",
        ],
    )
    add_para(doc, "Headline outcomes:", bold=True)
    add_bullets(
        doc,
        [
            "False-cross improved a lot: short 96% → 46%; fair long 88% → 38% (reliability held).",
            "Image-swap improved only modestly: about +7 percentage points short; almost flat on fair long.",
            "Honest claim: we fix routing leakage / anti-leakage grounding — not token-swap SOTA.",
        ],
    )

    # 2 Setup
    add_heading(doc, "2. System setup", 1)
    add_table(
        doc,
        ["Component", "Role"],
        [
            ["Backbone", "Frozen BLIP2-OPT-2.7B"],
            ["Editor", "LiveEdit — lifelong MoE pool of edit experts"],
            ["At edit time", "Write a new expert + keys into the pool"],
            ["At test time", "Retrieve relevant experts and inject residual edits into VLM layers"],
            ["Data", "VLKEB short / long grounding probes (image-swap, hard-negative, etc.)"],
        ],
    )
    add_para(
        doc,
        "Without our method, retrieval leans on compressed vision signals "
        "(Q-Former → few tokens → scalar similarity). That is brittle when many edits "
        "sit in one lifelong pool.",
    )

    # 3 Problem
    add_heading(doc, "3. Problem diagnosed", 1)
    add_para(doc, "Observation", bold=True)
    add_bullets(
        doc,
        [
            "On the correct image, reliability is high (~94% short).",
            "On a wrong / swapped image with a similar prompt, the model still often behaves as if the edit applies (~41% image-swap for LiveEdit baseline).",
            "Routing logs show false-cross ~96%: a foreign expert is retrieved.",
        ],
    )
    add_para(doc, "Why it matters", bold=True)
    add_para(
        doc,
        "Prior VLKE papers mostly report reliability and locality. They rarely measure "
        "adversarial image-swap / hard-negative behaviour, or which expert was retrieved. "
        "The failure mode was under-reported, but it is central for lifelong editors that "
        "accumulate many experts.",
    )
    add_para(doc, "Diagnosis (Steps 0–6)", bold=True)
    add_bullets(
        doc,
        [
            "Scalar Q-Former vision routing collapses fine visual identity.",
            "Offline DINOv2 can separate images well, but naïve fusion / decode gates / patch-only retrain did not fix end-to-end leakage.",
            "Leakage is dominated by wrong-expert retrieval (false-cross), not own-expert false-fire.",
        ],
    )
    add_para(
        doc,
        "Implication: we need (a) better training to bind vision correctly, and "
        "(b) a retrieval rule that requires vision ∧ text agreement.",
        italic=True,
    )

    # 4 Sample + methodology
    add_heading(doc, "4. Methodology with sample data (how and why)", 1)
    add_para(
        doc,
        "Naming used below: LiveEdit baseline (L0); AVBT (S2); AVBT + Dual-Expert Keys (DEK); "
        "Reliability (Rel); Image-swap (Swap); False-cross (FC).",
        size=10,
        italic=True,
    )

    add_heading(doc, "4.1 Running example — Sample A", 2)
    add_para(
        doc,
        "We walk one real VLKEB grounding sample end-to-end. Source: grounding_eval_short.json, index 0.",
    )
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Original prompt", "What city in Illinois is depicted in the image?"],
            ["Rephrase", "Which Illinois city is shown in the picture?"],
            ["Old model belief", "Laredo, Texas"],
            ["Edit target", "Kankakee"],
            [
                "Conflict prompt",
                "(The photograph is often misread as depicting Laredo, Texas.) What city in Illinois…?",
            ],
        ],
    )
    add_para(doc, "Probe images (left → right): edit / image-swap / hard-negative / modality-consistency")
    add_image(
        doc,
        SAMPLES / "sampleA_montage.jpg",
        width=6.2,
        caption="Figure 1. Sample A probe montage.",
    )

    add_heading(doc, "4.2 Step 0–1 — Edit write into the lifelong pool", 2)
    add_para(
        doc,
        "How: LiveEdit does not fine-tune the whole VLM. It builds a MoE expert e_Kankakee "
        "that later injects residuals so the answer becomes “Kankakee”, and stores retrieval keys "
        "in the lifelong pool.",
    )
    add_para(
        doc,
        "Why: lifelong editing needs cheap, additive updates. The risk is that many experts "
        "collide under weak vision keys.",
    )
    add_image(
        doc,
        SAMPLES / "sampleA_edit_kankakee.jpg",
        width=3.4,
        caption="Figure 2. Correct edit image used to build the vision key (expect “Kankakee” under Reliability).",
    )

    add_heading(doc, "4.3 Step 2 — AVBT: how and why on this sample", 2)
    add_para(
        doc,
        "How: AVBT adds an adversarial binding objective on top of the edit loss.",
        bold=True,
    )
    add_para(doc, "Positive pair (bind): same prompt family + correct image → strengthen e_Kankakee.")
    add_image(doc, SAMPLES / "sampleA_edit_kankakee.jpg", width=3.0, caption="Figure 3a. AVBT positive image.")
    add_para(doc, "Negative pair (suppress): same prompt + wrong city image → penalize firing e_Kankakee.")
    add_image(doc, SAMPLES / "sampleA_swap.jpg", width=3.0, caption="Figure 3b. AVBT negative — image-swap.")
    add_para(doc, "Harder negative (confusable partner):")
    add_image(doc, SAMPLES / "sampleA_hardneg.jpg", width=3.0, caption="Figure 3c. AVBT negative — hard-negative.")
    add_para(
        doc,
        "Conceptually: L = L_edit(correct → “Kankakee”) + λ · L_positive(correct → own expert) "
        "+ λ · L_negative(swap / hard-neg → suppress wrong experts).",
        italic=True,
    )
    add_para(
        doc,
        "Why: the editor must learn when not to bind, not only how to write the new string. "
        "In practice AVBT lifts image-swap (~41% → ~48%) but alone leaves false-cross ~96% — "
        "so training helps tokens a bit, while retrieval can still leak.",
    )

    add_heading(doc, "4.4 Step 3–6 — Dual-Expert Keys: how and why", 2)
    add_para(
        doc,
        "How: each expert stores separate vision keys (Q-Former + DINOv2 patch) and a text/prompt key. "
        "At every later query we keep an expert only if vision_ok ∧ text_ok. If either fails, the expert is blocked.",
    )
    add_para(doc, "Reliability query (correct image) — dual-gate should PASS → answer “Kankakee”:", bold=True)
    add_image(doc, SAMPLES / "sampleA_edit_kankakee.jpg", width=3.0, caption="Figure 4. Reliability query image.")
    add_para(
        doc,
        "Image-swap without Dual-Expert Keys (the bug): wrong city + Illinois-city prompt. "
        "Weak vision routing often still retrieves e_Kankakee → False-cross → may leak “Kankakee”.",
    )
    add_image(doc, SAMPLES / "sampleA_swap.jpg", width=3.0, caption="Figure 5. Image-swap query (leakage case).")
    add_para(
        doc,
        "Image-swap with Dual-Expert Keys (our fix): same query, but vision∧text agreement is required. "
        "For e_Kankakee the vision gate typically fails / disagrees on the non-Kankakee photo → expert blocked → "
        "false-cross drops (96% → 46% short).",
        bold=False,
    )
    add_para(
        doc,
        "Why DEK works: false-cross is a retrieval agreement failure. Requiring both modalities to agree "
        "removes many “looks similar but wrong entity” fires without rewriting the whole VLM.",
        italic=True,
    )
    add_para(doc, "Same logic applies to hard-negative; modality-consistency (same entity, new photo) should still pass:")
    add_image(doc, SAMPLES / "sampleA_modality.jpg", width=3.0, caption="Figure 6. Modality-consistency image (still expect “Kankakee”).")

    add_heading(doc, "4.5 Method summary on this sample", 2)
    add_table(
        doc,
        ["Stage", "What happens", "Our method’s role"],
        [
            ["Edit write", "Create e_Kankakee + keys in pool", "Store vision and text keys (DEK)"],
            ["AVBT +", "Bind own expert on correct image", "Positive binding loss"],
            ["AVBT −", "Suppress fires on swap / hard-neg", "Adversarial binding loss"],
            ["Rel test", "Retrieve e_Kankakee → “Kankakee”", "Dual-gate allows match"],
            ["Swap test", "Baseline: false-cross leak", "Dual-gate blocks mismatch"],
        ],
    )
    arch = DOC / "THESIS_architecture_AVBT_DEK.png"
    if arch.is_file():
        add_image(doc, arch, width=6.0, caption="Figure 7. Method architecture (AVBT + Dual-Expert Keys).")

    # 5 Results
    add_heading(doc, "5. Result analysis", 1)
    add_para(
        doc,
        "Protocol unless noted: dsn=100, sen=10, BLIP2-OPT-2.7B, grounding eval JSONs, meeting step 6.",
        size=10,
        italic=True,
    )

    add_heading(doc, "5.1 Component ablation (short, seed 42)", 2)
    add_para(
        doc,
        "Orthogonal effects: AVBT mainly moves token/adversarial training; Dual-Expert Keys mainly fixes retrieval.",
    )
    add_table(
        doc,
        ["Method", "Rel", "Swap", "FC", "Role"],
        [
            ["LiveEdit baseline (L0)", "93.5%", "40.7%", "96%", "baseline"],
            ["AVBT (S2)", "93.8%", "47.8%", "96%", "+swap, FC unchanged"],
            ["AVBT + Dual-Expert Keys (DEK)", "93.7%", "47.7%", "46%", "ours"],
            ["False-cross mining (FX)", "93.7%", "47.2%", "43%", "negative (no swap↑)"],
            ["Pool buffer (PB)", "93.7%", "48.0%", "55%", "negative (FC↑)"],
        ],
    )
    if (FIG / "aaai_ablation_short.png").is_file():
        add_image(doc, FIG / "aaai_ablation_short.png", width=5.8, caption="Figure 8. Short component ablation.")

    add_heading(doc, "5.2 Multi-seed confirmation (n=5)", 2)
    add_table(
        doc,
        ["Metric", "LiveEdit baseline (L0)", "Dual-Expert Keys (DEK)", "Δ", "p (paired t)"],
        [
            ["Image-swap", "41.2%", "48.6%", "+7.4 pp", "≈ 3.4×10⁻⁴"],
            ["False-cross", "96%", "46%", "−50 pp", "stable every seed"],
            ["Reliability", "93.5%", "93.7%", "+0.25 pp", "held"],
        ],
    )
    if (FIG / "aaai_multiseed_dek.png").is_file():
        add_image(doc, FIG / "aaai_multiseed_dek.png", width=5.8, caption="Figure 9. Multi-seed DEK vs LiveEdit baseline.")

    add_heading(doc, "5.3 Competitors (short)", 2)
    add_para(
        doc,
        "Fine-tuning / MEND-VL can look strong on swap only because reliability falls (~71–73%). "
        "Dual-Expert Keys keeps ~94% Rel and is the LiveEdit-family row that cuts false-cross. "
        "FC is n/a for FT/MEND/SERAC (no same routing pool).",
    )
    add_table(
        doc,
        ["Method", "Rel", "Swap", "Swap-gap", "FC"],
        [
            ["LiveEdit baseline (L0)", "93.5%", "40.7%", "52.8 pp", "96%"],
            ["AVBT + Dual-Expert Keys (DEK)", "93.7%", "47.7%", "46.0 pp", "46%"],
            ["Fine-tuning (FT)", "73.0%", "73.1%", "≈0", "n/a"],
            ["MEND-VL", "71.3%", "69.4%", "~2 pp", "n/a"],
            ["SERAC-VL", "38.6%", "38.6%", "≈0", "n/a"],
        ],
    )
    if (FIG / "aaai_competitors_short.png").is_file():
        add_image(doc, FIG / "aaai_competitors_short.png", width=5.8, caption="Figure 10. Short competitor comparison.")

    add_heading(doc, "5.4 Fair long-form (matched training)", 2)
    add_table(
        doc,
        ["Method", "Rel", "Swap", "FC"],
        [
            ["LiveEdit baseline (L0) long", "74.0%", "42.4%", "88%"],
            ["AVBT + Dual-Expert Keys (DEK) long", "72.8%", "42.7%", "38%"],
        ],
    )
    add_para(
        doc,
        "Same qualitative pattern as short: FC win + Rel held; swap almost flat. "
        "We do not claim long swap ≥55%. Short→long without retrain collapses Rel and is not used as headline long evidence.",
    )
    if (FIG / "aaai_long_fair.png").is_file():
        add_image(doc, FIG / "aaai_long_fair.png", width=5.2, caption="Figure 11. Fair long-form results.")

    add_heading(doc, "5.5 What failed (kept as negatives)", 2)
    add_bullets(
        doc,
        [
            "λ / data-size sweeps: no swap breakthrough.",
            "False-cross mining (FX): FC 46→43; swap flat.",
            "Pool buffer (PB): gate FAIL (swap 48%, FC 55%).",
            "Stretch goal long swap ≥55%: FAIL.",
        ],
    )
    if (FIG / "aaai_results_panel.png").is_file():
        add_image(doc, FIG / "aaai_results_panel.png", width=6.0, caption="Figure 12. Results panel summary.")

    # 6 Positioning
    add_heading(doc, "6. Positioning vs LiveEdit / AnyEdit / publication", 1)
    add_bullets(
        doc,
        [
            "vs LiveEdit (CVPR 2025): we extend their stack; novelty is the grounding-leakage / false-cross problem + AVBT/DEK fix — not a from-scratch lifelong VLM editor.",
            "vs AnyEdit (ICML 2025): orthogonal (text LLM long-form editing). Cite only; do not claim to beat it.",
            "Thesis: strong if framed as anti-leakage routing with Rel preserved.",
            "AAAI main: stretch — do not frame as “LiveEdit++”; need Gen/Loc parity table + probe audit for a stronger submission case.",
        ],
    )

    add_heading(doc, "7. Recommended claim wording", 1)
    add_para(
        doc,
        "Say: Lifelong VLM editors leak under adversarial images via false-cross expert retrieval. "
        "AVBT improves adversarial binding; Dual-Expert Keys require vision∧text agreement at retrieval. "
        "Together they cut false-cross roughly in half while preserving reliability on short (multi-seed) and fair long-form VLKEB.",
        italic=True,
    )
    add_para(
        doc,
        "Do not say: We achieve state-of-the-art image-swap on long-form / we beat LiveEdit and AnyEdit on all benchmarks.",
        bold=True,
    )

    add_heading(doc, "8. Next steps discussed with supervisor", 1)
    add_bullets(
        doc,
        [
            "Confirm framing: problem+mechanism (FC) vs LiveEdit-extension language.",
            "Extract clean Rel / T-Gen / M-Gen / T-Loc / M-Loc side-by-side for L0 vs DEK.",
            "Complete human audit of 150 swap/hard-neg probes (sheet ready).",
            "Decide thesis-first vs AAAI stretch timeline.",
        ],
    )

    add_para(
        doc,
        "Bottom line: clear problem, mechanistic diagnosis, dual-gate method that works on the failure mode it targets, "
        "multi-seed evidence, competitor context, and honest negatives. The win is anti-leakage routing (false-cross) "
        "with reliability preserved — not a dramatic long-form swap leap.",
        bold=True,
    )

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
