"""Generate thesis presentation PPTX from structured content."""
from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "THESIS_PRESENTATION.pptx"

# Colors (RGB tuples for reference in shapes if needed)
TITLE_COLOR = None  # default theme


def _set_title(slide, text: str, subtitle: str | None = None):
    slide.shapes.title.text = text
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def _add_bullets(slide, lines: list[str], left=0.5, top=1.4, width=9.0, height=5.5, font_size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        # strip markdown bold markers for display
        line = line.replace("**", "")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(font_size)
        if line.startswith("- "):
            p.text = line[2:]
            p.level = 0


def _add_table(slide, headers: list[str], rows: list[list[str]], top=1.5, col_widths=None):
    nrows = len(rows) + 1
    ncols = len(headers)
    left, width, height = Inches(0.4), Inches(9.2), Inches(0.35 * nrows)
    table = slide.shapes.add_table(nrows, ncols, Inches(0.4), Inches(top), Inches(9.2), height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(11)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val).replace("**", "")
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)


def _add_flow_text(slide, title: str, lines: list[str], top=1.3):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(0.4))
    box.text_frame.paragraphs[0].text = title
    box.text_frame.paragraphs[0].font.bold = True
    box.text_frame.paragraphs[0].font.size = Pt(14)
    _add_bullets(slide, lines, top=top + 0.35, font_size=14)


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    # Slide 1 — Title
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = "Autoregressive Chunk-Wise Editing for\nLong-Form Visual Knowledge Editing"
    s.placeholders[1].text = (
        "Extending LiveEdit on VLKEB-long (BLIP2-OPT-2.7b)\n"
        "MSc Thesis — Progress Presentation\n"
        "[Your Name] · [Institution] · June 2026"
    )

    # Slide 2 — Introduction
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Introduction — Motivation")
    _add_bullets(s, [
        "Vision–language models encode outdated or incorrect facts about the world.",
        "Model editing updates specific knowledge without full retraining.",
        "LiveEdit (CVPR 2025): lifelong VLM editing with low-rank mixture-of-experts on VLKEB.",
        "Gap: real knowledge is often long-form (65+ tokens); single-pass editors struggle.",
        "Our work: VLKEB-long dataset + AR-LiveEdit (chunk-wise autoregressive editing).",
    ], top=1.3)

    # Slide 3 — Problem
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Problem Statement")
    _add_bullets(s, [
        "LiveEdit edits the full target in one pass — hard for ~89-token targets.",
        "Efficacy barrier: single-token / single-pass edits weakly affect distant tokens (AnyEdit).",
        "Exact match (EM) ≈ 0% on long answers; token- and chunk-level metrics matter.",
        "Sequential lifelong editing (50 edits/batch) may stress locality.",
        "",
        "RQ1: Does AR-LiveEdit improve long-form editing vs baseline while preserving generality & locality?",
        "RQ2: How does inference chunk size (8/16/32/64) affect performance?",
    ], top=1.2, font_size=16)

    # Slide 4 — Objectives
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Objectives & Contributions")
    _add_table(s,
        ["#", "Objective", "Status"],
        [
            ["1", "Build VLKEB-long (anchored long alt)", "Done"],
            ["2", "Train baseline + AR-LiveEdit (20 epochs)", "Done"],
            ["3", "Full eval ~3150 edits (seed 42)", "Done"],
            ["4", "Compare baseline vs AR", "Done"],
            ["5", "Chunk size ablation 8/16/32/64", "Done"],
        ],
        top=1.35, col_widths=[0.5, 6.5, 1.5])
    _add_bullets(s, [
        "Contributions: VLKEB-long + Ollama pipeline; AR vs baseline study; chunk ablation analysis.",
    ], top=4.2, font_size=14)

    # Slide 5 — Lit review editing
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review — Model Editing")
    _add_bullets(s, [
        "ROME, MEMIT, MEND: locate-then-edit or hypernetwork-based LLM editing.",
        "Standard metrics: Efficacy/Reliability, Generalization, Specificity/Locality.",
        "Limitation: single-token or short-span edits → efficacy barrier for 65+ token targets.",
    ], top=1.3)

    # Slide 6 — LiveEdit
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review — LiveEdit (CVPR 2025)")
    _add_bullets(s, [
        "Lifelong vision–language model editing with MoE experts per edit.",
        "Hard routing: filter visually irrelevant experts.",
        "Soft routing: fuse textually relevant experts.",
        "Training: reliability + generality + KL locality losses.",
        "VLKEB: Rel., T-Gen., M-Gen., T-Loc., M-Loc. — but short targets only.",
    ], top=1.3)

    # Slide 7 — AnyEdit
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Literature Review — AnyEdit (ICML 2025)")
    _add_bullets(s, [
        "Autoregressive chunk-wise editing for long-form LLM knowledge.",
        "Chain rule of mutual information motivates chunk decomposition.",
        "Metrics: ROUGE-L, BERTScore for long outputs.",
        "We adapt chunk-wise editing to LiveEdit VLM + lifelong setting → AR-LiveEdit.",
    ], top=1.3)

    # Slide 8 — Gap
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Research Gap & Positioning")
    _add_table(s,
        ["Work", "Modality", "Lifelong", "Long targets"],
        [
            ["ROME/MEMIT/MEND", "LLM", "Limited", "Poor"],
            ["AnyEdit", "LLM", "Batch", "Strong"],
            ["LiveEdit", "VLM", "Yes", "Short VLKEB"],
            ["This thesis", "VLM", "Yes (seq=50)", "VLKEB-long"],
        ],
        top=1.4, col_widths=[2.2, 1.3, 1.5, 2.0])
    _add_bullets(s, [
        "Position: lifelong VLM editing + long-form targets + chunk-wise AR.",
    ], top=4.5, font_size=14)

    # Slide 9 — Methodology overview
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Methodology — Pipeline Overview")
    _add_flow_text(s, "End-to-end flow:", [
        "VLKEB (official) → Ollama long-form generator → VLKEB-long JSON",
        "Train LiveEdit on VLKEB-long: Baseline (no AR) + AR (chunk_size=16)",
        "Backbone: BLIP2-OPT-2.7b (frozen) + trainable editor modules",
        "Evaluate: sequential_edit_n=50, seed=42, ~3150 edits",
        "Metrics: Reliability, Generality, Locality, Chunk EM/F1",
    ])

    # Slide 10 — Dataset
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Methodology — VLKEB-long Dataset")
    _add_bullets(s, [
        "Source: official VLKEB train/eval + images (unchanged).",
        "alt_short = original short alt; alt = long-form via Ollama (gemma4:31b-cloud).",
        "Validation: alt_short substring in long alt; token length band (BLIP2 tokenizer).",
        "Scale: 3174 eval rows; 3150 edits in full eval.",
        "Limitation: lexical anchoring ≠ full visual grounding of every phrase.",
    ], top=1.25, font_size=16)

    # Slide 11 — AR method
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Methodology — LiveEdit vs AR-LiveEdit")
    _add_bullets(s, [
        "Baseline: one edit_one_piece call on full target_new.",
        "AR-LiveEdit: split target into token chunks (size 16) → edit prefix chunk-by-chunk.",
        "At inference: expert pool routes per chunk (chunk-aware routing gate).",
        "Training: ar_mode + per-chunk loss aggregation.",
        "Example: ~89 tokens → ~6 chunk edits at cs=16 vs 1 edit for baseline.",
    ], top=1.25, font_size=16)

    # Slide 12 — Eval protocol
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Methodology — Evaluation Protocol")
    _add_table(s,
        ["Metric", "Measures"],
        [
            ["Reliability (acc)", "Token accuracy on edit → long target"],
            ["Text / image generality", "Rephrased prompt / image"],
            ["Text / image locality", "Unrelated QA unchanged"],
            ["Chunk EM / F1", "Per-chunk exact match & token F1"],
            ["65+ bin", "3149 of 3150 samples (~89 tokens mean)"],
        ],
        top=1.35, col_widths=[2.5, 6.5])
    _add_bullets(s, [
        "Baseline ckpt: epoch-20 ema_loss 0.79 | AR ckpt: epoch-20 ema_loss 1.51",
    ], top=4.6, font_size=13)

    # Slide 13 — Main results
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Results — Baseline vs AR (cs=16)")
    _add_table(s,
        ["Metric", "Baseline", "AR", "Δ"],
        [
            ["Reliability acc", "75.40%", "75.94%", "+0.54 pp"],
            ["Chunk EM", "13.50%", "14.82%", "+1.32 pp"],
            ["Chunk F1", "76.28%", "76.90%", "+0.62 pp"],
            ["Text rephrase", "75.21%", "75.63%", "+0.42 pp"],
            ["Image rephrase", "71.49%", "72.51%", "+1.02 pp"],
            ["Locality (T/I)", "100%/100%", "100%/100%", "0"],
            ["Exact match", "0%", "0%", "—"],
        ],
        top=1.3, col_widths=[2.4, 1.5, 1.5, 1.3])
    _add_bullets(s, [
        "Headline: AR improves long-form editing without locality drop.",
        "Validation: long-form improvement ✓ | locality bounded ✓",
    ], top=4.8, font_size=14)

    # Slide 14 — Chunk ablation
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Results — Chunk Size Ablation (AR)")
    _add_table(s,
        ["chunk_size", "Rel acc", "Chunk EM", "Chunk F1", "Locality"],
        [
            ["8", "75.94%", "26.26%", "76.48%", "100%"],
            ["16 (default)", "75.94%", "14.82%", "76.90%", "100%"],
            ["32", "75.94%", "8.11%", "77.30%", "100%"],
            ["64", "75.94%", "6.14%", "77.88%", "100%"],
        ],
        top=1.35, col_widths=[1.5, 1.3, 1.3, 1.3, 1.2])
    _add_bullets(s, [
        "Token accuracy identical across chunk sizes.",
        "cs=16 matches training; recommended default.",
        "Chunk EM not comparable across sizes (different windows).",
    ], top=4.5, font_size=14)

    # Slide 15 — Discussion
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Discussion")
    _add_bullets(s, [
        "AR gives consistent gains: +0.54 pp acc, +1.32 pp chunk EM.",
        "Generality up; locality remains 100% on this run.",
        "Gains are modest but directionally consistent on new long-form split.",
        "EM = 0% expected for ~89-token targets.",
        "Inference chunk size does not change token acc when trained at cs=16.",
        "Limits: single backbone; Ollama long alt not fully visually verified.",
    ], top=1.2, font_size=16)

    # Slide 16 — Future work
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Limitations & Future Work")
    _add_table(s,
        ["Limitation", "Future direction"],
        [
            ["Small absolute gains", "Multiple seeds, significance tests"],
            ["Dataset grounding", "Human audit + grounding tiers"],
            ["Metrics", "ROUGE-L, BERTScore"],
            ["Novelty", "Vision-gated chunks, forgetting analysis"],
            ["Single VLLM", "Second backbone (LLaVA)"],
            ["Lifelong depth", "Edit-count sweep 1→1000"],
        ],
        top=1.35, col_widths=[3.5, 5.5])

    # Slide 17 — Conclusion
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Conclusion")
    _add_bullets(s, [
        "Built VLKEB-long with documented Ollama protocol.",
        "Trained & evaluated baseline LiveEdit vs AR-LiveEdit at full scale (3150 edits).",
        "AR beats baseline on reliability, chunk EM, generality — no locality loss.",
        "chunk_size=16 is the sensible default (matches training).",
        "",
        "Take-home: AR chunk-wise editing is a viable extension of LiveEdit for long-form VLM editing.",
    ], top=1.25, font_size=17)

    # Slide 18 — Q&A
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Thank You — Questions?")
    _add_bullets(s, [
        "Backup: LiveEdit CVPR 2025 architecture figure",
        "Results: eval_results/comparison/*.json",
        "",
        "Anticipated Q: Why small improvement? → hard task, long outputs, single seed.",
        "Anticipated Q: VLKEB-long novelty? → extension + protocol, not new images.",
        "Anticipated Q: Why not EM? → ~89 tokens; use token acc & chunk metrics.",
    ], top=1.4, font_size=16)

    # Optional: architecture text slide
    s = prs.slides.add_slide(content_layout)
    _set_title(s, "Backup — System Architecture")
    _add_flow_text(s, "Components:", [
        "Input: image + prompt + long target alt",
        "BLIP2-OPT-2.7b (frozen): vision encoder + OPT decoder",
        "LiveEdit editor: expert generator → expert pool",
        "Hard visual routing + soft text routing → residual update to hidden states",
        "Output: edited VLM prediction on reliability / generality / locality probes",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
