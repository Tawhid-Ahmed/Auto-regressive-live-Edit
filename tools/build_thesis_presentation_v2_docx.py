"""Generate thesis presentation DOCX (V2) from doc/THESIS_PRESENTATION_V2.md.

Renders Mermaid diagram blocks to PNG via local @mermaid-js/mermaid-cli and
embeds them in the Word document.

Outputs doc/THESIS_PRESENTATION_V2.docx
"""
from __future__ import annotations

import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from math_docx_format import MATH_FONT, add_formatted_text, fill_table_cell

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRC = ROOT / "doc" / "THESIS_PRESENTATION_V2.md"
DEFAULT_OUT = ROOT / "doc" / "THESIS_PRESENTATION_V2.docx"
DEFAULT_IMG_DIR = ROOT / "doc" / "_presentation_v2_diagrams"
MMDC = (
    ROOT
    / "tools"
    / "mermaid-cli"
    / "node_modules"
    / ".bin"
    / ("mmdc.cmd" if sys.platform == "win32" else "mmdc")
)


def _render_mermaid(code: str, out_png: Path) -> bool:
    out_png.parent.mkdir(parents=True, exist_ok=True)
    mmd = out_png.with_suffix(".mmd")
    mmd.write_text(code.strip() + "\n", encoding="utf-8")
    cmd = [
        str(MMDC),
        "-i",
        str(mmd),
        "-o",
        str(out_png),
        "-b",
        "transparent",
        "-w",
        "1400",
    ]
    try:
        subprocess.run(cmd, check=True, capture_output=True, text=True)
    except (subprocess.CalledProcessError, FileNotFoundError) as exc:
        print(f"Warning: failed to render {out_png.name}: {exc}", file=sys.stderr)
        return False
    return out_png.exists()


def _add_rich_paragraph(doc: Document, text: str, *, bullet: bool = False, size: int = 11):
    text = text.strip()
    if not text:
        return
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    add_formatted_text(p, text, size=size, font=MATH_FONT)


def _add_equation_line(doc: Document, text: str):
    """Render display equations with Cambria Math and sub/superscript runs."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_formatted_text(p, text.strip(), size=12, font=MATH_FONT, equation=True)


def _parse_table_lines(lines: list[str]) -> tuple[list[str], list[list[str]]] | None:
    if len(lines) < 2:
        return None
    rows = [line.strip() for line in lines if line.strip().startswith("|")]
    if len(rows) < 2:
        return None
    headers = [c.strip() for c in rows[0].strip("|").split("|")]
    body_rows: list[list[str]] = []
    for row in rows[2:]:
        body_rows.append([c.strip().replace("**", "") for c in row.strip("|").split("|")])
    return headers, body_rows


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, header in enumerate(headers):
        fill_table_cell(table.rows[0].cells[j], header.replace("**", ""), bold=True, size=10)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            fill_table_cell(table.rows[i].cells[j], str(val), size=10)
    doc.add_paragraph()


def _split_slides(text: str) -> list[str]:
    parts = re.split(r"\n---\n", text)
    slides: list[str] = []
    for part in parts:
        part = part.strip()
        if not part or part.startswith("# How to use"):
            continue
        if part.startswith("# Thesis Presentation"):
            continue
        slides.append(part)
    return slides


def _process_slide(doc: Document, slide_text: str, diagram_idx: list[int]):
    lines = slide_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            doc.add_heading(title, level=1)
            i += 1
            continue

        img_match = re.match(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
        if img_match:
            rel = img_match.group(1).replace("/", os.sep)
            img_path = (ROOT / "doc" / rel) if not Path(rel).is_absolute() else Path(rel)
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_picture(str(img_path), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                _add_rich_paragraph(doc, f"[Image not found: {rel}]")
            i += 1
            continue

        if stripped.startswith("```mermaid"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # closing ```
            code = "\n".join(block)
            diagram_idx[0] += 1
            png = IMG_DIR / f"diagram_{diagram_idx[0]:02d}.png"
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if _render_mermaid(code, png):
                doc.add_picture(str(png), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                cap = caption.add_run(f"[Diagram {diagram_idx[0]} — render failed; see Mermaid in source MD]")
                cap.italic = True
                cap.font.size = Pt(9)
            continue

        if stripped.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            parsed = _parse_table_lines(table_lines)
            if parsed:
                _add_table(doc, *parsed)
            continue

        if line.startswith("    ") and stripped:
            _add_equation_line(doc, stripped)
            i += 1
            continue

        if stripped.startswith(">"):
            quote = stripped.lstrip("> ").strip()
            p = doc.add_paragraph(quote)
            p.paragraph_format.left_indent = Inches(0.25)
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(10)
            i += 1
            continue

        if stripped.startswith("- "):
            _add_rich_paragraph(doc, stripped[2:], bullet=True)
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            _add_rich_paragraph(doc, stripped, bullet=False)
            i += 1
            continue

        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph(stripped.strip("*"))
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(10)
            i += 1
            continue

        if stripped:
            _add_rich_paragraph(doc, stripped)
        i += 1


def build(
    src: Path = DEFAULT_SRC,
    out: Path = DEFAULT_OUT,
    img_dir: Path = DEFAULT_IMG_DIR,
):
    if not src.exists():
        raise FileNotFoundError(src)
    if not MMDC.exists():
        raise FileNotFoundError(
            f"Mermaid CLI not found at {MMDC}. Run: "
            "npm install @mermaid-js/mermaid-cli --prefix tools/mermaid-cli"
        )

    global IMG_DIR
    IMG_DIR = img_dir

    text = src.read_text(encoding="utf-8")
    doc = Document()
    doc.core_properties.title = out.stem
    doc.core_properties.subject = "Autoregressive Chunk-Wise Editing for Long-Form Knowledge in VLMs"

    doc.add_heading(
        "Autoregressive Chunk-Wise Editing for Long-Form Knowledge in Vision–Language Models",
        level=0,
    )
    sub = doc.add_paragraph(
        f"Thesis presentation — generated from {src.name} with embedded diagrams."
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    diagram_idx = [0]
    for slide in _split_slides(text):
        _process_slide(doc, slide, diagram_idx)
        doc.add_page_break()

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Wrote {out} ({diagram_idx[0]} diagrams)")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Build thesis presentation DOCX from markdown.")
    parser.add_argument("--src", type=Path, default=DEFAULT_SRC)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--img-dir", type=Path, default=DEFAULT_IMG_DIR)
    args = parser.parse_args()
    build(args.src, args.out, args.img_dir)
